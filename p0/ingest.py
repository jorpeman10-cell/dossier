# -*- coding: utf-8 -*-
"""ingest.py — turn a resume file (txt/md/docx/pdf) into plain text.

Zero third-party deps for txt/md; docx and pdf are used only if the
optional libs are installed (python-docx / pypdf / PyMuPDF).
"""
from __future__ import annotations
import sys
from pathlib import Path


def ingest(path: str | Path) -> str:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    suffix = p.suffix.lower()
    if suffix in (".txt", ".md", ".markdown"):
        return p.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        return _docx(p)
    if suffix == ".pdf":
        return _pdf(p)
    # last resort: try utf-8 text
    return p.read_text(encoding="utf-8", errors="replace")


def _docx(p: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        raise RuntimeError("python-docx not installed; save as txt/pdf instead")
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    for table in d.tables:  # resumes often use tables for basic info
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(x for x in parts if x.strip())


def _pdf(p: Path) -> str:
    try:
        import fitz  # PyMuPDF, best Chinese support
        doc = fitz.open(str(p))
        return "\n".join(page.get_text() for page in doc)
    except ImportError:
        pass
    try:
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(p)).pages)
    except ImportError:
        raise RuntimeError("neither PyMuPDF nor pypdf installed; save as txt instead")


if __name__ == "__main__":
    print(ingest(sys.argv[1]))
